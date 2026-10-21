"""Focused regression tests for the Attendance four-metric update.

Run: python -m unittest test_attendance_single_punch_pct.py
Uses only synthetic figures; no employee records or real reports are bundled.
"""
from __future__ import annotations

import ast
import io
import sqlite3
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

import db
import parsers

ROOT = Path(__file__).resolve().parent.parent


def record(day: str, site: str, metric: str, value: float, section: str = "attendance") -> dict:
    return {
        "section": section, "data_date": day,
        "site_key": db.normalize_site(site), "site": site,
        "facility_type": "XD" if site.endswith(" XD") else "DC",
        "metric": metric, "value": value, "unit": "count",
        "source_file": "synthetic.xlsx",
    }


class AttendanceSinglePunchPercentTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.path = Path(self.tempdir.name) / "attendance.db"
        db.init_db(self.path)

    def get_pct(self, day: str, site: str):
        with sqlite3.connect(self.path) as conn:
            row = conn.execute(
                "SELECT value,unit FROM summary_metrics WHERE section='attendance' "
                "AND data_date=? AND site_key=? AND metric='Single Punch %'",
                (day, db.normalize_site(site)),
            ).fetchone()
        return row

    def test_exact_site_date_percentage_and_original_count_preserved(self):
        db.replace_summary_rows(self.path, [
            record("2026-10-01", "NCR Bamnoli DC", "Active HC", 320),
            record("2026-10-01", "NCR Bamnoli DC", "Single Punch", 9),
            record("2026-10-02", "NCR Bamnoli DC", "Active HC", 200),
            record("2026-10-02", "NCR Bamnoli DC", "Single Punch", 10),
            record("2026-10-01", "Hubli XD", "Active HC", 40),
            record("2026-10-01", "Hubli XD", "Single Punch", 1),
            record("2026-10-01", "NCR Bamnoli DC", "Active HC", 17, "roster"),
        ])
        db.sync_attendance_single_punch_percent(self.path)
        self.assertAlmostEqual(self.get_pct("2026-10-01", "NCR Bamnoli DC")[0], 2.8125, places=8)
        self.assertEqual(self.get_pct("2026-10-01", "NCR Bamnoli DC")[1], "percent")
        self.assertAlmostEqual(self.get_pct("2026-10-02", "NCR Bamnoli DC")[0], 5.0)
        self.assertAlmostEqual(self.get_pct("2026-10-01", "Hubli XD")[0], 2.5)
        with sqlite3.connect(self.path) as conn:
            self.assertEqual(conn.execute(
                "SELECT value FROM summary_metrics WHERE section='attendance' AND data_date='2026-10-01' "
                "AND site_key='ncr bamnoli dc' AND metric='Single Punch'"
            ).fetchone()[0], 9)
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM summary_metrics WHERE metric='Single Punch %' AND section='roster'"
            ).fetchone()[0], 0)

    def test_zero_missing_and_zero_single_punch(self):
        db.replace_summary_rows(self.path, [
            record("2026-10-01", "Zero HC DC", "Active HC", 0),
            record("2026-10-01", "Zero HC DC", "Single Punch", 0),
            record("2026-10-01", "Missing HC DC", "Single Punch", 5),
            record("2026-10-01", "Zero Punch DC", "Active HC", 15),
            record("2026-10-01", "Zero Punch DC", "Single Punch", 0),
        ])
        db.sync_attendance_single_punch_percent(self.path)
        self.assertIsNone(self.get_pct("2026-10-01", "Zero HC DC"))
        self.assertIsNone(self.get_pct("2026-10-01", "Missing HC DC"))
        self.assertEqual(self.get_pct("2026-10-01", "Zero Punch DC"), (0.0, "percent"))

    def test_reimports_upsert_and_remove_stale_invalid_denominator(self):
        db.replace_summary_rows(self.path, [
            record("2026-10-01", "A DC", "Active HC", 100),
            record("2026-10-01", "A DC", "Single Punch", 5),
            record("2026-10-02", "A DC", "Active HC", 20),
            record("2026-10-02", "A DC", "Single Punch", 3),
        ])
        db.sync_attendance_single_punch_percent(self.path)
        self.assertEqual(self.get_pct("2026-10-01", "A DC")[0], 5)
        db.replace_summary_rows(self.path, [
            record("2026-10-01", "A DC", "Active HC", 50),
            record("2026-10-01", "A DC", "Single Punch", 8),
        ])
        db.sync_attendance_single_punch_percent(self.path, ["2026-10-01"])
        self.assertEqual(self.get_pct("2026-10-01", "A DC")[0], 16)
        self.assertEqual(self.get_pct("2026-10-02", "A DC")[0], 15)
        db.sync_attendance_single_punch_percent(self.path)
        with sqlite3.connect(self.path) as conn:
            self.assertEqual(conn.execute(
                "SELECT COUNT(*) FROM summary_metrics WHERE section='attendance' AND metric='Single Punch %'"
            ).fetchone()[0], 2)
        db.replace_summary_rows(self.path, [record("2026-10-01", "A DC", "Active HC", 0)])
        db.sync_attendance_single_punch_percent(self.path, ["2026-10-01"])
        self.assertIsNone(self.get_pct("2026-10-01", "A DC"))
        self.assertEqual(self.get_pct("2026-10-02", "A DC")[0], 15)

    def test_full_attendance_import_creates_percentage_and_keeps_sp_drilldown(self):
        wb = Workbook()
        dc = wb.active
        dc.title = "Attendance Adherence - DC"
        dc.append(["Site", datetime(2026, 10, 1), None, None])
        dc.append([None, "Active HC", "Single Punch", "8 Hrs + Adherence"])
        dc.append(["Example DC", 80, 4, 0.75])
        xd = wb.create_sheet("Attendance Adherence - XD")
        xd.append(["Site", datetime(2026, 10, 1), None, None])
        xd.append([None, "Active HC", "Single Punch", "8 Hrs + Adherence"])
        base = wb.create_sheet("Attendance Base")
        base.append([
            "To be considered", "To be considered(Values - TRUE/FALSE)",
            "user_id", "user_name", "role", "facility_name",
            datetime(2026, 10, 1), None, None, None,
        ])
        base.append([None, None, None, None, None, None, "shift", "total_hours", "working_hours", "punch_status"])
        base.append(["Yes", False, "Emp-1", "Sample Worker", "Picker", "Example DC", "A", 8, 8, "SP"])
        buffer = io.BytesIO()
        wb.save(buffer)
        wb.close()
        result = parsers.process_workbook(buffer.getvalue(), "sample_attendance.xlsx", str(self.path))
        self.assertEqual(result["report_type"], "Attendance")
        self.assertEqual(result["detail_rows"], 1)
        self.assertEqual(self.get_pct("2026-10-01", "Example DC"), (5.0, "percent"))
        with sqlite3.connect(self.path) as conn:
            self.assertEqual(conn.execute(
                "SELECT value FROM summary_metrics WHERE section='attendance' "
                "AND site_key='example dc' AND metric='8 Hrs + Adherence'"
            ).fetchone()[0], 75.0)
            self.assertEqual(conn.execute(
                "SELECT punch_status FROM attendance_detail WHERE employee_id='Emp-1'"
            ).fetchone()[0], "SP")

    def test_metric_options_and_employee_drill_down_mapping_are_correct(self):
        # Inspect app.py without executing Streamlit UI side effects.
        source = (ROOT / "app.py").read_text(encoding="utf-8")
        tree = ast.parse(source)
        attendance = next(node for node in tree.body
                          if isinstance(node, ast.FunctionDef) and node.name == "render_attendance")
        option_node = next(node for node in ast.walk(attendance)
                           if isinstance(node, ast.Assign)
                           and any(isinstance(t, ast.Name) and t.id == "attendance_options" for t in node.targets))
        options = ast.literal_eval(option_node.value)
        self.assertEqual(options, [
            "8 Hrs + Adherence", "5 to 8 Hrs Adherence",
            "<5 Hrs Adherence", "Single Punch %",
        ])
        self.assertIn('"Single Punch %": "Single Punch"', source)
        self.assertIn('"Single Punch %": "SP"', source)


if __name__ == "__main__":
    unittest.main()
