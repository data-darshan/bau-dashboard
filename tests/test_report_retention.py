"""Storage/import contract: no historical employee data on disk or across sessions."""
from __future__ import annotations

import io
import sqlite3
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

import db
from parsers import process_workbook
from site_palette import ensure_site_colors, site_color


def make_roster(*, day: int, site: str, count: int) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title in ("DC Summary", "XD Summary", "DC Roster Adoption", "XD Roster Adoption"):
        ws = wb.create_sheet(title)
        ws.append(["Site", datetime(2026, 10, day), None, None])
        ws.append([None, "Active HC", "Rostered HC", "Roster Match"])
        if title == "DC Summary" and site.endswith(" DC"):
            ws.append([site, count, count - 1, count - 2])
        if title == "XD Summary" and site.endswith(" XD"):
            ws.append([site, count, count - 1, count - 2])
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def make_onboarding() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title in ("Aadhar - DC", "Aadhar - XD", "Hyperverge - DC", "Hyperverge - XD"):
        ws = wb.create_sheet(title)
        ws.append(["Site", datetime(2026, 10, 13), None])
        if title.startswith("Aadhar"):
            ws.append([None, "Aadhar", "Manual"])
            if title.endswith(" DC"):
                ws.append(["Bangalore DC", 1, 1])
        else:
            ws.append([None, "RED Cases"])
            if title.endswith(" DC"):
                ws.append(["Bangalore DC", 1])
    base = wb.create_sheet("Onboarding Base")
    base.append(["ID", "Name", "Role", "Registered Facility", "Onboarded Date",
                 "Approved By", "Onboarded by Aadhaar (Y/N)", "Hyperverge Status",
                 "Auto Verification Status"])
    base.append(["A-1", "Employee One", "Picker", "Bangalore DC", datetime(2026, 10, 13),
                 "HR A", "Y", "red", "Pending"])
    base.append(["A-2", "Employee Two", "Packer", "Bangalore DC", datetime(2026, 10, 13),
                 "HR B", "N", "green", "Verified"])
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def make_flex() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title in ("DC Summary - 1", "XD Summary - 1", "DC Summary - 2", "XD Summary - 2"):
        ws = wb.create_sheet(title)
        if title.endswith("- 1"):
            metrics = ["Total Flex Used", "Onboarded & Punched", "Onboarded & Not Punched", "Not Onboarded"]
            row = [2, 0, 1, 1]
        else:
            metrics = ["Total Punched", "5 to 8 Hours", "Less than 5 Hours", "Single Punch", "8 Hours+", "8 Hrs+ Adherence"]
            row = [1, 1, 0, 0, 0, 0.0]
        ws.append(["Site", datetime(2026, 10, 13)] + [None] * (len(metrics) - 1))
        ws.append([None] + metrics)
        if title.startswith("DC"):
            ws.append(["Hyderabad DC"] + row)
    base = wb.create_sheet("Central Flex Base")
    base.append(["Name", "Location", "Date", "HRBP Name", "Employee Type",
                 "Final Validation", "Valinor ID", "Check", "Flex Check", "Working Hours", "Bucket",
                 "Mobile No", "Aadhar Number", "Site HR"])
    base.append(["Employee One", "Hyderabad REV DC", datetime(2026, 10, 13), "BP One", "FT",
                 "onboarded", "V-1", "HF", "absent", 6.5, "5 to 8 hours", "888800001", "9999000001", "HR A"])
    base.append(["Employee Two", "Hyderabad DC", datetime(2026, 10, 13), "BP Two", "FT",
                 "not onboarded", "V-2", "-", "absent", 0, "-", "888800002", "9999000002", "HR B"])
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


class SessionRetentionTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.settings_db = Path(self.folder.name) / "site_settings.db"
        db.init_site_registry(self.settings_db)
        db.upsert_sites(self.settings_db, [("Bangalore DC", "DC", "South")])
        self.uri, self.anchor = db.new_session_database()
        self.addCleanup(self.anchor.close)
        db.init_db(self.uri)
        db.refresh_sites_from_registry(self.uri, self.settings_db)

    def count(self, table: str) -> int:
        with db.connect(self.uri) as con:
            return con.execute(f"SELECT count(*) FROM {table}").fetchone()[0]

    def test_successive_uploads_replace_only_their_own_report_domain(self):
        a = process_workbook(make_roster(day=1, site="Bangalore DC", count=14), "older.xlsx", self.uri)
        self.assertEqual(a["report_type"], "Roster")
        self.assertEqual(self.count("summary_metrics"), 3)
        b = process_workbook(make_roster(day=13, site="Bangalore DC", count=42), "newer.xlsx", self.uri)
        self.assertEqual(b["report_type"], "Roster")
        self.assertEqual(self.count("summary_metrics"), 3)
        self.assertEqual(self.count("import_history"), 1)
        with db.connect(self.uri) as con:
            dates = {r[0] for r in con.execute("SELECT DISTINCT data_date FROM summary_metrics")}
            self.assertEqual(dates, {"2026-10-13"})
            self.assertEqual(con.execute("SELECT source_file FROM summary_metrics LIMIT 1").fetchone()[0], "newer.xlsx")

    def test_import_progress_reports_actual_stages_without_persisting_history(self):
        events = []
        result = process_workbook(
            make_roster(day=13, site="Bangalore DC", count=42),
            "progress.xlsx", self.uri,
            progress=lambda fraction, caption: events.append((fraction, caption)),
        )
        self.assertEqual(result["report_type"], "Roster")
        self.assertGreaterEqual(len(events), 5)
        self.assertEqual(events[0][1], "Opening Excel workbook")
        self.assertTrue(any("Parsing Roster" in caption for _, caption in events))
        self.assertTrue(any("Replacing older Roster" in caption for _, caption in events))
        self.assertEqual(events[-1][0], 0.87)
        self.assertEqual(self.count("summary_metrics"), 3)

    def test_bad_new_report_does_not_delete_working_current_data(self):
        process_workbook(make_roster(day=1, site="Bangalore DC", count=14), "working.xlsx", self.uri)
        book = Workbook()
        book.remove(book.active)
        for name in ("DC Summary", "XD Summary", "DC Roster Adoption", "XD Roster Adoption"):
            ws = book.create_sheet(name)
            ws.append(["Site", datetime(2026, 10, 13)])
            ws.append([None, "Active HC"])
        blob = io.BytesIO()
        book.save(blob)
        book.close()
        with self.assertRaisesRegex(ValueError, "no usable summary"):
            process_workbook(blob.getvalue(), "empty.xlsx", self.uri)
        self.assertEqual(self.count("summary_metrics"), 3)
        with db.connect(self.uri) as con:
            self.assertEqual(con.execute("SELECT DISTINCT data_date FROM summary_metrics").fetchone()[0], "2026-10-01")

    def test_import_transaction_rolls_back_after_failed_insert(self):
        process_workbook(make_roster(day=1, site="Bangalore DC", count=14), "working.xlsx", self.uri)
        bad = {"section": "roster", "data_date": "2026-10-12", "site_key": None,
               "site": "Broken", "facility_type": "DC", "metric": "Active HC", "unit": "count", "value": 5}
        with self.assertRaises(sqlite3.IntegrityError):
            with db.connect(self.uri) as con:
                con.execute("BEGIN IMMEDIATE")
                db.clear_previous_report(con, "Roster")
                db.replace_summary_rows(self.uri, [bad], con=con)
        self.assertEqual(self.count("summary_metrics"), 3)
        self.assertEqual(self.count("import_history"), 1)

    def test_session_isolation_and_durable_registry(self):
        process_workbook(make_roster(day=13, site="Bangalore DC", count=42), "private.xlsx", self.uri)
        new_uri, new_anchor = db.new_session_database()
        try:
            db.init_db(new_uri)
            db.refresh_sites_from_registry(new_uri, self.settings_db)
            with db.connect(new_uri) as con:
                self.assertEqual(con.execute("SELECT count(*) FROM summary_metrics").fetchone()[0], 0)
                self.assertEqual(con.execute("SELECT region FROM site_mapping WHERE site_key='bangalore dc'").fetchone()[0], "South")
        finally:
            new_anchor.close()
        files = {p.name for p in Path(self.folder.name).iterdir() if p.is_file()}
        self.assertEqual(files, {"site_settings.db"})
        with sqlite3.connect(self.settings_db) as con:
            tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        self.assertEqual(tables, {"site_mapping"})

    def test_new_site_mappings_and_colors_persist_not_personal_data(self):
        db.upsert_sites(self.uri, [("New XD", "XD", None)])
        db.save_discovered_sites_to_registry(self.uri, self.settings_db)
        db.refresh_sites_from_registry(self.uri, self.settings_db)
        saved = ensure_site_colors(self.settings_db)
        self.assertIn("new xd", saved)
        self.assertEqual(site_color("new xd", db_path=self.settings_db), saved["new xd"])
        with sqlite3.connect(self.settings_db) as con:
            tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            self.assertEqual(tables, {"site_mapping", "site_color_mapping"})
            self.assertEqual(con.execute("SELECT region FROM site_mapping WHERE site_key='new xd'").fetchone()[0], "UNASSIGNED")

    def test_all_domains_and_onboarding_hr_breakdown_in_memory(self):
        from onboarding_breakdown import import_hr_site_counts
        process_workbook(make_roster(day=1, site="Bangalore DC", count=14), "roster.xlsx", self.uri)
        onb_blob = make_onboarding()
        onb = process_workbook(onb_blob, "onboarding.xlsx", self.uri)
        self.assertEqual(onb["report_type"], "Onboarding")
        counts = import_hr_site_counts(onb_blob)
        self.assertEqual(len(counts), 2)
        self.assertEqual({(r["hr"], r["method"], r["count"]) for r in counts},
                         {("HR A", "Aadhar", 1), ("HR B", "Manual", 1)})
        flex = process_workbook(make_flex(), "flex.xlsx", self.uri)
        self.assertEqual(flex["report_type"], "Flex")
        with db.connect(self.uri) as con:
            sections = {r[0] for r in con.execute("SELECT DISTINCT section FROM summary_metrics")}
            self.assertEqual(sections, {"roster", "onboarding_aadhar", "onboarding_hyperverge", "flex_usage", "flex_working"})
            self.assertEqual(con.execute("SELECT count(*) FROM onboarding_red_detail").fetchone()[0], 1)
            self.assertEqual(con.execute("SELECT count(*) FROM flex_detail").fetchone()[0], 2)
            fields = con.execute("SELECT phone_number,aadhaar_number,hr_name FROM flex_detail WHERE employee_id='V-1'").fetchone()
            self.assertEqual(tuple(fields), ("888800001", "9999000001", "HR A"))
        process_workbook(make_roster(day=13, site="Bangalore DC", count=40), "roster_latest.xlsx", self.uri)
        with db.connect(self.uri) as con:
            self.assertEqual(con.execute("SELECT count(*) FROM flex_detail").fetchone()[0], 2)
            self.assertEqual(con.execute("SELECT count(*) FROM onboarding_red_detail").fetchone()[0], 1)
            self.assertEqual(con.execute("SELECT count(*) FROM import_history").fetchone()[0], 3)

    def test_flex_out_of_scope_base_dates_and_sites_never_enter_session(self):
        raw = make_flex()
        book = __import__('openpyxl').load_workbook(io.BytesIO(raw))
        ws = book['Central Flex Base']
        good = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
        outside_month = list(good); outside_month[2] = datetime(2026, 11, 10)
        next_day = list(good); next_day[2] = datetime(2026, 10, 14)
        other_site = list(good); other_site[1] = 'Unrelated DC'
        ws.append(outside_month); ws.append(next_day); ws.append(other_site)
        buffer = io.BytesIO();book.save(buffer);book.close()
        result = process_workbook(buffer.getvalue(), 'dirty-flex.xlsx', self.uri)
        self.assertEqual((result['min_date'], result['max_date']), ('2026-10-13','2026-10-13'))
        self.assertEqual(result['detail_rows'], 2)
        self.assertEqual(self.count('flex_detail'), 2)
        with db.connect(self.uri) as con:
            dates = {r[0] for r in con.execute('SELECT DISTINCT data_date FROM flex_detail')}
            self.assertEqual(dates, {'2026-10-13'})
            self.assertEqual(con.execute("SELECT COUNT(*) FROM flex_detail WHERE site_key='hyderabad rev dc'").fetchone()[0], 1)

    def test_explicit_clear_erases_session_reports_but_not_sites(self):
        process_workbook(make_roster(day=13, site="Bangalore DC", count=42), "private.xlsx", self.uri)
        db.clear_session_reports(self.uri)
        self.assertEqual(self.count("summary_metrics"), 0)
        self.assertEqual(self.count("import_history"), 0)
        self.assertEqual(self.count("site_mapping"), 1)


if __name__ == "__main__":
    unittest.main()
