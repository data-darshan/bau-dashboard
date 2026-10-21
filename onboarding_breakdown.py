"""Pure in-memory Aadhaar/Manual HR-by-site breakdown for the latest upload.

The official chart totals still come from workbook summary sheets. Employee
names and workbook bytes are never saved to disk by this module.
"""
from __future__ import annotations

import io
from collections import Counter
from datetime import date
from typing import Any, Sequence

import polars as pl
from openpyxl import load_workbook

from db import normalize_site
from parsers import clean_header, parse_summary_sheet, to_iso_date

METHODS = ("Aadhar", "Manual")
REQUIRED_SHEETS = {"Aadhar - DC", "Aadhar - XD", "Onboarding Base"}
REQUIRED_HEADERS = (
    "Registered Facility", "Onboarded Date", "Approved By", "Onboarded by Aadhaar (Y/N)",
)


def _method(flag: Any) -> str | None:
    value = clean_header(flag).casefold()
    if value in {"y", "yes", "true"}:
        return "Aadhar"
    if value in {"n", "no", "false"}:
        return "Manual"
    return None


def extract_hr_site_counts(blob: bytes) -> tuple[list[dict], set[tuple[str, str, str]]]:
    """Extract summaries for the exact date/site/method cells shown by the donut.

    The Onboarding Base can include test/other facility categories and days
    beyond the official DC/XD summary; ignore them rather than inflating totals.
    """
    book = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        missing_sheets = REQUIRED_SHEETS.difference(book.sheetnames)
        if missing_sheets:
            raise ValueError(f"Onboarding workbook missing worksheets: {', '.join(sorted(missing_sheets))}")

        summary = []
        for sheet, facility_type in (("Aadhar - DC", "DC"), ("Aadhar - XD", "XD")):
            summary.extend(parse_summary_sheet(
                book[sheet], "onboarding_aadhar", facility_type, "", {"Compliance %"}, "fraction",
            ))

        expected: dict[tuple[str, str, str], tuple[str, str, int]] = {}
        for metric_row in summary:
            if metric_row["metric"] not in METHODS:
                continue
            metric_count = float(metric_row["value"])
            if metric_count < 0 or metric_count != int(metric_count):
                raise ValueError("Expected a non-negative whole-number onboarding count")
            key = (metric_row["data_date"], metric_row["site_key"], metric_row["metric"])
            expected[key] = (metric_row["site"], metric_row["facility_type"], int(metric_count))

        ws = book["Onboarding Base"]
        header = [clean_header(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        index = {name: pos for pos, name in enumerate(header) if name}
        missing_headers = [name for name in REQUIRED_HEADERS if name not in index]
        if missing_headers:
            raise ValueError(f"Onboarding Base missing required headers: {', '.join(missing_headers)}")

        counts: Counter[tuple[str, str, str, str]] = Counter()
        site_col = index["Registered Facility"]
        date_col = index["Onboarded Date"]
        hr_col = index["Approved By"]
        method_col = index["Onboarded by Aadhaar (Y/N)"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            day = to_iso_date(row[date_col])
            if not day:
                continue
            method = _method(row[method_col])
            if not method:
                continue
            site_key = normalize_site(clean_header(row[site_col]))
            if (day, site_key, method) not in expected:
                continue
            hr = clean_header(row[hr_col])
            if not hr or hr == "-":
                hr = "Not specified"
            counts[(day, site_key, method, hr)] += 1

        actual: Counter[tuple[str, str, str]] = Counter()
        for (day, site_key, method, _), count in counts.items():
            actual[(day, site_key, method)] += count
        mismatches = [
            (key, expected[key][2], actual[key])
            for key in expected if expected[key][2] != actual[key]
        ]
        if mismatches:
            sample = "; ".join(
                f"{d}/{site}/{m}: summary {target}, HR detail {got}"
                for (d, site, m), target, got in mismatches[:3]
            )
            raise ValueError(
                "Onboarding HR/site counts do not reconcile with official DC/XD summaries. "
                f"No breakdown was saved. {len(mismatches)} mismatch(es); {sample}"
            )

        records = []
        for (day, site_key, method, hr), count in sorted(counts.items()):
            site_name, facility_type, _ = expected[(day, site_key, method)]
            records.append({
                "date": day, "site_key": site_key, "site": site_name,
                "facility_type": facility_type, "hr": hr, "method": method,
                "count": count,
            })
        return records, set(expected)
    finally:
        book.close()


def import_hr_site_counts(blob: bytes) -> list[dict]:
    """Validate the workbook and return ephemeral HR/site counts.

    The Streamlit session holds these records only until another Onboarding
    report replaces them or the browser session closes.
    """
    records, _covered = extract_hr_site_counts(blob)
    return records


def query_hr_site_counts(
    records: list[dict],
    start: date,
    end: date,
    method: str,
    regions: Sequence[str],
    facility_types: Sequence[str],
    site_keys: Sequence[str],
    mapping: pl.DataFrame,
) -> pl.DataFrame:
    """Filter this session's most recently uploaded breakdown.

    Region mapping is resolved at display time from the existing mapping data;
    changing a region in Settings therefore does not require a second upload.
    """
    cols = ["Site HR", "Site", "Count"]
    if method not in METHODS:
        raise ValueError("Unknown onboarding method")
    selected = []
    region_by_site = dict(zip(mapping["site_key"], mapping["Region"]))
    display_by_site = dict(zip(mapping["site_key"], mapping["Site"]))
    begin, finish = start.isoformat(), end.isoformat()
    for rec in records:
        if rec["method"] != method or not begin <= rec["date"] <= finish:
            continue
        key = rec["site_key"]
        if regions and region_by_site.get(key, "UNASSIGNED") not in regions:
            continue
        if facility_types and rec["facility_type"] not in facility_types:
            continue
        if site_keys and key not in site_keys:
            continue
        selected.append({
            "Site HR": rec["hr"],
            "Site": display_by_site.get(key, rec["site"]),
            "site_key": key,
            "Count": int(rec["count"]),
        })
    if not selected:
        return pl.DataFrame(schema={"Site HR": pl.String, "Site": pl.String, "Count": pl.Int64})
    frame = pl.DataFrame(selected)
    # Include site_key to prevent merging different facilities with the same label.
    grouped = (
        frame.group_by(["site_key", "Site", "Site HR"])
        .agg(pl.col("Count").sum())
        .sort(["Count", "Site", "Site HR"], descending=[True, False, False])
    )
    return grouped.select(cols)
