from __future__ import annotations

import gc
import io
import re
from datetime import date, datetime
from typing import Any, Callable

from openpyxl import load_workbook

from db import (
    clear_previous_report,
    connect,
    add_import_history,
    infer_facility_type,
    normalize_site,
    replace_attendance_detail,
    replace_flex_detail,
    replace_onboarding_red_detail,
    replace_summary_rows,
    sync_attendance_single_punch_percent,
    upsert_sites,
)


def clean_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\n", " ").strip())


def to_iso_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value in (None, "", "-"):
        return None
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d %b %Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def to_float(value: Any) -> float | None:
    if value in (None, "", "-"):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def percent_value(value: Any, scale: str) -> float | None:
    v = to_float(value)
    if v is None:
        return None
    return v * 100.0 if scale == "fraction" else v


def parse_summary_sheet(ws, section: str, facility_type: str, source_file: str, percent_metrics: set[str], percent_scale: str) -> list[dict]:
    row1 = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    row2 = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    columns: list[tuple[int, str, str]] = []
    current_date: str | None = None
    for idx in range(1, max(len(row1), len(row2))):
        h1 = row1[idx] if idx < len(row1) else None
        h2 = row2[idx] if idx < len(row2) else None
        maybe_date = to_iso_date(h1)
        if maybe_date:
            current_date = maybe_date
        metric = clean_header(h2)
        if current_date and metric:
            columns.append((idx, current_date, metric))

    out: list[dict] = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        site = clean_header(row[0] if row else None)
        if not site:
            continue
        if normalize_site(site) in {"grand total", "total", "overall total"}:
            continue
        site_key = normalize_site(site)
        for idx, data_date, metric in columns:
            if idx >= len(row):
                continue
            raw = row[idx]
            if raw is None:
                continue
            unit = "percent" if metric in percent_metrics else "count"
            value = percent_value(raw, percent_scale) if unit == "percent" else to_float(raw)
            if value is None:
                continue
            out.append({
                "section": section,
                "data_date": data_date,
                "site_key": site_key,
                "site": site,
                "facility_type": facility_type,
                "metric": metric,
                "value": value,
                "unit": unit,
                "source_file": source_file,
            })
    return out


def detect_report_type(sheet_names: list[str]) -> str:
    s = set(sheet_names)
    if {"Attendance Adherence - DC", "Attendance Adherence - XD", "Attendance Base"}.issubset(s):
        return "Attendance"
    if {"Aadhar - DC", "Aadhar - XD", "Hyperverge - DC", "Hyperverge - XD", "Onboarding Base"}.issubset(s):
        return "Onboarding"
    if {"DC Summary - 1", "DC Summary - 2", "XD Summary - 1", "XD Summary - 2", "Central Flex Base"}.issubset(s):
        return "Flex"
    if {"DC Summary", "XD Summary", "DC Roster Adoption", "XD Roster Adoption"}.issubset(s):
        return "Roster"
    raise ValueError("This workbook does not match the expected Attendance, Onboarding, Flex, or Roster report structure.")


def _site_rows_from_summary(rows: list[dict]) -> list[tuple[str, str, None]]:
    seen = {}
    for r in rows:
        seen[r["site_key"]] = (r["site"], r["facility_type"], None)
    return list(seen.values())


def parse_attendance(wb, source_file: str) -> tuple[list[dict], list[dict], list[str]]:
    pct = {"Punched %", "Double Punch Adherence", "<5 Hrs Adherence", "5 to 8 Hrs Adherence", "8 Hrs + Adherence"}
    summary = []
    summary += parse_summary_sheet(wb["Attendance Adherence - DC"], "attendance", "DC", source_file, pct, "fraction")
    summary += parse_summary_sheet(wb["Attendance Adherence - XD"], "attendance", "XD", source_file, pct, "fraction")

    ws = wb["Attendance Base"]
    row1 = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    row2 = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    header_idx = {clean_header(v): i for i, v in enumerate(row1) if clean_header(v)}
    required = ["To be considered", "To be considered(Values - TRUE/FALSE)", "user_id", "user_name", "role", "facility_name"]
    missing = [h for h in required if h not in header_idx]
    if missing:
        raise ValueError(f"Attendance Base is missing required column(s): {', '.join(missing)}")

    date_blocks: list[dict] = []
    current_date = None
    block: dict[str, Any] = {}
    for idx, (top, sub) in enumerate(zip(row1, row2)):
        maybe_date = to_iso_date(top)
        if maybe_date:
            current_date = maybe_date
            block = {"date": current_date}
            date_blocks.append(block)
        subh = clean_header(sub).lower()
        if current_date and subh in {"shift", "total_hours", "working_hours", "punch_status"}:
            block[subh] = idx

    detail: list[dict] = []
    # Attendance Base includes the next day's partial records; its dates must
    # not extend the reporting range beyond the official summary sheets.
    report_pairs = {(r["data_date"], r["site_key"]) for r in summary}
    dates = sorted({r["data_date"] for r in summary})
    date_blocks = [b for b in date_blocks if b["date"] in set(dates)]
    yes_i = header_idx["To be considered"]
    false_i = header_idx["To be considered(Values - TRUE/FALSE)"]
    id_i = header_idx["user_id"]
    name_i = header_idx["user_name"]
    role_i = header_idx["role"]
    site_i = header_idx["facility_name"]

    for row in ws.iter_rows(min_row=3, values_only=True):
        if clean_header(row[yes_i]).casefold() != "yes":
            continue
        false_value = row[false_i]
        is_false = false_value is False or clean_header(false_value).casefold() == "false"
        if not is_false:
            continue
        employee_id = clean_header(row[id_i])
        site = clean_header(row[site_i])
        if not employee_id or not site:
            continue
        site_key = normalize_site(site)
        for b in date_blocks:
            if (b["date"], site_key) not in report_pairs:
                continue
            pidx = b.get("punch_status")
            if pidx is None or pidx >= len(row):
                continue
            status = clean_header(row[pidx]).upper()
            # Dashboard employee drill-down is needed only for the exception groups.
            if status not in {"SP", "MIS", "HF"}:
                continue
            detail.append({
                "data_date": b["date"],
                "site_key": site_key,
                "site": site,
                "employee_id": employee_id,
                "employee_name": clean_header(row[name_i]),
                "role": clean_header(row[role_i]),
                "shift": clean_header(row[b.get("shift")]) if b.get("shift") is not None else None,
                "total_hours": to_float(row[b.get("total_hours")]) if b.get("total_hours") is not None else None,
                "working_hours": to_float(row[b.get("working_hours")]) if b.get("working_hours") is not None else None,
                "punch_status": status,
                "source_file": source_file,
            })
    return summary, detail, dates


def parse_onboarding(wb, source_file: str) -> tuple[list[dict], list[dict], list[str]]:
    aadhar_pct = {"Compliance %"}
    summary = []
    summary += parse_summary_sheet(wb["Aadhar - DC"], "onboarding_aadhar", "DC", source_file, aadhar_pct, "fraction")
    summary += parse_summary_sheet(wb["Aadhar - XD"], "onboarding_aadhar", "XD", source_file, aadhar_pct, "fraction")
    summary += parse_summary_sheet(wb["Hyperverge - DC"], "onboarding_hyperverge", "DC", source_file, set(), "fraction")
    summary += parse_summary_sheet(wb["Hyperverge - XD"], "onboarding_hyperverge", "XD", source_file, set(), "fraction")

    ws = wb["Onboarding Base"]
    headers = [clean_header(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers) if h}
    required = ["ID", "Name", "Role", "Registered Facility", "Onboarded Date", "Approved By", "Onboarded by Aadhaar (Y/N)", "Hyperverge Status", "Auto Verification Status"]
    missing = [h for h in required if h not in idx]
    if missing:
        raise ValueError(f"Onboarding Base is missing required column(s): {', '.join(missing)}")

    detail = []
    summary_pairs = {(r["data_date"], r["site_key"]) for r in summary}
    dates = sorted({r["data_date"] for r in summary})
    for row in ws.iter_rows(min_row=2, values_only=True):
        data_date = to_iso_date(row[idx["Onboarded Date"]])
        hv = clean_header(row[idx["Hyperverge Status"]]).casefold()
        if not data_date or hv != "red":
            continue
        site = clean_header(row[idx["Registered Facility"]])
        if (data_date, normalize_site(site)) not in summary_pairs:
            continue
        detail.append({
            "data_date": data_date,
            "site_key": normalize_site(site),
            "site": site,
            "employee_id": clean_header(row[idx["ID"]]),
            "employee_name": clean_header(row[idx["Name"]]),
            "role": clean_header(row[idx["Role"]]),
            "approved_by": clean_header(row[idx["Approved By"]]),
            "aadhaar_onboarding": clean_header(row[idx["Onboarded by Aadhaar (Y/N)"]]),
            "hyperverge_status": clean_header(row[idx["Hyperverge Status"]]),
            "auto_verification_status": clean_header(row[idx["Auto Verification Status"]]),
            "comments": clean_header(row[idx.get("Auto Verification Comments")]) if idx.get("Auto Verification Comments") is not None else None,
            "source_file": source_file,
        })
    return summary, detail, dates


def parse_flex(wb, source_file: str) -> tuple[list[dict], list[dict], list[str]]:
    usage_pct = {"Compliance %"}
    working_pct = {"Double Punch Adherence", "8 Hrs+ Adherence"}
    summary = []
    summary += parse_summary_sheet(wb["DC Summary - 1"], "flex_usage", "DC", source_file, usage_pct, "fraction")
    summary += parse_summary_sheet(wb["XD Summary - 1"], "flex_usage", "XD", source_file, usage_pct, "fraction")
    summary += parse_summary_sheet(wb["DC Summary - 2"], "flex_working", "DC", source_file, working_pct, "fraction")
    summary += parse_summary_sheet(wb["XD Summary - 2"], "flex_working", "XD", source_file, working_pct, "fraction")

    ws = wb["Central Flex Base"]
    headers = [clean_header(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers) if h}

    def optional_index(*names: str) -> int | None:
        """Return an optional Central Flex Base column, tolerant of punctuation."""
        def canonical(value: str) -> str:
            return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())

        folded = {canonical(name): pos for name, pos in idx.items()}
        for name in names:
            found = folded.get(canonical(name))
            if found is not None:
                return found
        return None

    phone_i = optional_index(
        "Phone Number", "Phone No", "Mobile Number", "Mobile No", "Mobile",
        "Phone", "Contact Number", "Contact No",
    )
    aadhaar_i = optional_index(
        "Aadhaar Number", "Aadhar Number", "Aadhaar No", "Aadhar No",
        "Aadhaar", "Aadhar", "Aadhaar Card Number", "Aadhar Card Number",
    )
    hr_i = optional_index("HR Name", "Site HR", "Site HR Name", "HR", "HRR", "HRR Name", "Respective HR")
    required = ["Name", "Location", "Date", "HRBP Name", "Employee Type", "Final Validation", "Valinor ID", "Check", "Flex Check", "Working Hours", "Bucket"]
    missing = [h for h in required if h not in idx]
    if missing:
        raise ValueError(f"Central Flex Base is missing required column(s): {', '.join(missing)}")

    detail = []
    report_pairs = {(r["data_date"], r["site_key"]) for r in summary}
    dates = sorted({r["data_date"] for r in summary})
    for row in ws.iter_rows(min_row=2, values_only=True):
        data_date = to_iso_date(row[idx["Date"]])
        site = clean_header(row[idx["Location"]])
        if not data_date or not site:
            continue
        site_key = normalize_site(site)
        # The supplied Oct 17 Central Flex Base includes date cells in
        # January-September/November-December and the next unsummarized day.
        # Do not re-interpret source dates speculatively; exclude rows whose
        # date/site has no corresponding official DC/XD summary. REV DC rows
        # legitimately roll into their parent DC summary.
        report_key = site_key[:-7] + " dc" if site_key.endswith(" rev dc") else site_key
        if (data_date, site_key) not in report_pairs and (data_date, report_key) not in report_pairs:
            continue
        detail.append({
            "data_date": data_date,
            "site_key": site_key,
            "site": site,
            "employee_id": clean_header(row[idx["Valinor ID"]]),
            "employee_name": clean_header(row[idx["Name"]]),
            "role": None,
            "hrbp_name": clean_header(row[idx["HRBP Name"]]),
            "employee_type": clean_header(row[idx["Employee Type"]]),
            "validation_status": clean_header(row[idx["Final Validation"]]),
            "flex_check": clean_header(row[idx["Flex Check"]]),
            "punch_code": clean_header(row[idx["Check"]]),
            "working_hours": to_float(row[idx["Working Hours"]]),
            "bucket": clean_header(row[idx["Bucket"]]),
            "phone_number": clean_header(row[phone_i]) if phone_i is not None and phone_i < len(row) else None,
            "aadhaar_number": clean_header(row[aadhaar_i]) if aadhaar_i is not None and aadhaar_i < len(row) else None,
            "hr_name": clean_header(row[hr_i]) if hr_i is not None and hr_i < len(row) else None,
            "source_file": source_file,
        })
    return summary, detail, dates


def parse_roster(wb, source_file: str) -> tuple[list[dict], list[str]]:
    pct = {"Roster Adoption", "Roster Adherence"}
    summary = []
    summary += parse_summary_sheet(wb["DC Summary"], "roster", "DC", source_file, pct, "points")
    summary += parse_summary_sheet(wb["XD Summary"], "roster", "XD", source_file, pct, "points")
    dates = sorted({r["data_date"] for r in summary})
    return summary, dates


def process_workbook(
    blob: bytes,
    source_file: str,
    db_path: str,
    progress: Callable[[float, str], None] | None = None,
) -> dict:
    """Parse and atomically replace one report, emitting genuine stage milestones.

    Progress values reflect completed STAGES, not invented row throughput.
    Existing callers can omit the callback without changing import behavior.
    """
    def report_progress(value: float, caption: str) -> None:
        if progress is not None:
            progress(value, caption)

    report_progress(0.06, "Opening Excel workbook")
    bio = io.BytesIO(blob)
    wb = None
    try:
        wb = load_workbook(bio, read_only=True, data_only=True)
        report_progress(0.17, "Identifying workbook type")
        report_type = detect_report_type(wb.sheetnames)
        report_progress(0.28, f"Parsing {report_type} summary and employee detail sheets")
        if report_type == "Attendance":
            summary, detail, dates = parse_attendance(wb, source_file)
        elif report_type == "Onboarding":
            summary, detail, dates = parse_onboarding(wb, source_file)
        elif report_type == "Flex":
            summary, detail, dates = parse_flex(wb, source_file)
        else:
            summary, dates = parse_roster(wb, source_file)
            detail = []

        report_progress(0.65, f"Validating {len(summary):,} summary rows and {len(detail):,} detail rows")
        if not summary:
            raise ValueError("Report contains no usable summary rows; previous upload is unchanged.")

        # Atomic latest-upload-only import: a failed insert cannot wipe the
        # previous snapshot, and older dates from another upload never linger.
        # Each of the four domains is independent within this browser session.
        report_progress(0.74, f"Replacing older {report_type} data in this session")
        with connect(db_path) as con:
            con.execute("BEGIN IMMEDIATE")
            clear_previous_report(con, report_type)
            upsert_sites(db_path, _site_rows_from_summary(summary), con=con)
            summary_count = replace_summary_rows(db_path, summary, con=con)
            if report_type == "Attendance":
                sync_attendance_single_punch_percent(
                    db_path, sorted({r["data_date"] for r in summary}), con=con
                )
                detail_count = replace_attendance_detail(db_path, detail, dates, con=con)
            elif report_type == "Onboarding":
                detail_count = replace_onboarding_red_detail(db_path, detail, dates, con=con)
            elif report_type == "Flex":
                detail_count = replace_flex_detail(db_path, detail, dates, con=con)
            else:
                detail_count = 0
            add_import_history(db_path, report_type, source_file, dates,
                               summary_count, detail_count, con=con)
        report_progress(0.87, f"{report_type} rows imported into session memory")
        return {
            "report_type": report_type,
            "source_file": source_file,
            "min_date": min(dates) if dates else None,
            "max_date": max(dates) if dates else None,
            "summary_rows": summary_count,
            "detail_rows": detail_count,
        }
    finally:
        if wb is not None:
            wb.close()
        bio.close()
        del wb
        del bio
        gc.collect()
